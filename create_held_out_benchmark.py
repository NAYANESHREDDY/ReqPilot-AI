"""
Creates an independent, held-out benchmark specification and gold reference for
Warehouse Logistics & Fleet Dispatch System (held_out_logistics).
This domain is completely unseen in the training/remediation cycle.
"""

import json
from pathlib import Path
import docx

bench_dir = Path("evaluation/multi_project/benchmark/held_out_logistics")
bench_dir.mkdir(parents=True, exist_ok=True)

# 1. Requirements list (20 requirements across 5 canonical capabilities)
requirements = [
    # AUTHENTICATION
    {
        "id": "FR-001",
        "statement": "Logistics operators and warehouse dispatchers should be able to log in securely using corporate multi-factor authentication.",
        "summary": "Logistics operators and dispatchers should log in securely.",
        "canonical": "AUTHENTICATION",
        "display": "Authentication & Authorization",
        "type": "functional",
        "rationale": "Operator login and credential verification."
    },
    {
        "id": "FR-002",
        "statement": "The system should enforce role-based access control restricting warehouse forklift drivers to assigned loading bays.",
        "summary": "Role-based access control for forklift drivers to assigned bays.",
        "canonical": "AUTHENTICATION",
        "display": "Authentication & Authorization",
        "type": "functional",
        "rationale": "RBAC and physical facility zone permissions."
    },
    {
        "id": "FR-003",
        "statement": "Supervisory credentials and manager override approval should be required before overriding dangerous cargo safety limits.",
        "summary": "Manager override approval required for dangerous cargo safety limits.",
        "canonical": "AUTHENTICATION",
        "display": "Authentication & Authorization",
        "type": "functional",
        "rationale": "Supervisory override and security authority control."
    },
    {
        "id": "FR-004",
        "statement": "System sessions should automatically expire after 15 minutes of terminal inactivity on ruggedized handheld scanners.",
        "summary": "Session timeout after 15 minutes of scanner inactivity.",
        "canonical": "AUTHENTICATION",
        "display": "Authentication & Authorization",
        "type": "functional",
        "rationale": "Terminal session token lifecycle management."
    },

    # BOOKING (Slot & Bay Scheduling, Vehicle Dispatch Booking)
    {
        "id": "FR-005",
        "statement": "Freight carriers should be able to book loading dock appointments and reserve freight arrival time slots.",
        "summary": "Carriers can book loading dock appointments and time slots.",
        "canonical": "BOOKING",
        "display": "Booking & Scheduling",
        "type": "functional",
        "rationale": "Dock appointment scheduling and slot reservation."
    },
    {
        "id": "FR-006",
        "statement": "Warehouse supervisors should be able to schedule outbound freight delivery routes and allocate trucks to delivery windows.",
        "summary": "Schedule delivery routes and allocate trucks to windows.",
        "canonical": "BOOKING",
        "display": "Booking & Scheduling",
        "type": "functional",
        "rationale": "Route delivery scheduling and resource slot reservation."
    },
    {
        "id": "FR-007",
        "statement": "The system should prevent a loading bay marked unavailable for maintenance from accepting dock reservations.",
        "summary": "Prevent unavailable loading bay from accepting dock reservations.",
        "canonical": "BOOKING",
        "display": "Booking & Scheduling",
        "type": "functional",
        "rationale": "Reservation constraint blocking bookings on unavailable bays."
    },
    {
        "id": "FR-008",
        "statement": "Dock appointments should automatically expire if the transport vehicle fails to check in within 30 minutes of scheduled arrival.",
        "summary": "Dock appointment expiration after 30 minutes late check-in.",
        "canonical": "BOOKING",
        "display": "Booking & Scheduling",
        "type": "functional",
        "rationale": "Booking expiry policy."
    },

    # NOTIFICATION
    {
        "id": "FR-009",
        "statement": "The system should send automated SMS alerts to freight drivers when their designated loading dock is ready.",
        "summary": "Send SMS alerts to drivers when loading dock is ready.",
        "canonical": "NOTIFICATION",
        "display": "Notification Service",
        "type": "functional",
        "rationale": "Automated driver SMS dispatch notification."
    },
    {
        "id": "FR-010",
        "statement": "Warehouse managers should receive instant push notifications when a temperature-controlled cold-storage container experiences a temperature anomaly.",
        "summary": "Push notifications for cold-storage temperature anomalies.",
        "canonical": "NOTIFICATION",
        "display": "Notification Service",
        "type": "functional",
        "rationale": "Dispatched operational alert notification."
    },
    {
        "id": "FR-011",
        "statement": "Automated email notifications should be dispatched to shippers when customs clearance paperwork is rejected.",
        "summary": "Email notification to shippers upon customs paperwork rejection.",
        "canonical": "NOTIFICATION",
        "display": "Notification Service",
        "type": "functional",
        "rationale": "Automated email event notification."
    },
    {
        "id": "FR-012",
        "statement": "The system should dispatch alert notifications to terminal staff when dangerous goods dwell time exceeds permissible safety limits.",
        "summary": "Alert terminal staff when dangerous goods dwell time exceeds limit.",
        "canonical": "NOTIFICATION",
        "display": "Notification Service",
        "type": "functional",
        "rationale": "Safety warning notification delivery."
    },

    # DATA (Application Data Store: Warehouse Inventory, Pallet Master Records, Audit Logs)
    {
        "id": "FR-013",
        "statement": "The system should maintain persistent inventory records of pallet barcode, SKU, aisle location, bin number, weight, and expiry date.",
        "summary": "Maintain persistent inventory records of pallet, SKU, and aisle location.",
        "canonical": "DATA",
        "display": "Application Data Store",
        "type": "functional",
        "rationale": "Core warehouse inventory master record storage."
    },
    {
        "id": "FR-014",
        "statement": "Immutable database audit logs should preserve every pallet movement and bin transfer for at least 7 years for regulatory compliance.",
        "summary": "Preserve pallet transfer audit logs for 7 years.",
        "canonical": "DATA",
        "display": "Application Data Store",
        "type": "functional",
        "rationale": "Compliance audit trail data persistence."
    },
    {
        "id": "FR-015",
        "statement": "The system should record forklift equipment telematics, battery charging cycles, and inspection history in the central database.",
        "summary": "Record forklift telematics and inspection history in database.",
        "canonical": "DATA",
        "display": "Application Data Store",
        "type": "functional",
        "rationale": "Fleet equipment data store records."
    },
    {
        "id": "FR-016",
        "statement": "Database backup and disaster recovery replication should ensure zero data loss for freight manifests across secondary storage clusters.",
        "summary": "Database backup and replication for freight manifests.",
        "canonical": "DATA",
        "display": "Application Data Store",
        "type": "functional",
        "rationale": "Data store backup and persistence durability."
    },

    # REPORTING & ANALYTICS
    {
        "id": "FR-017",
        "statement": "Operations executives should be able to generate quarterly freight throughput reports and warehouse capacity utilization analytics.",
        "summary": "Generate freight throughput reports and utilization analytics.",
        "canonical": "REPORTING",
        "display": "Reporting & Analytics",
        "type": "functional",
        "rationale": "Analytical report generation and business utilization metrics."
    },
    {
        "id": "FR-018",
        "statement": "The system should calculate average dock turnaround times and carrier on-time delivery statistical distributions.",
        "summary": "Calculate dock turnaround times and on-time delivery statistics.",
        "canonical": "REPORTING",
        "display": "Reporting & Analytics",
        "type": "functional",
        "rationale": "Operational KPI metrics and performance analytics."
    },
    {
        "id": "FR-019",
        "statement": "Logistics managers should be able to export carbon emission summaries and fleet fuel consumption reports.",
        "summary": "Export carbon emission summaries and fuel consumption reports.",
        "canonical": "REPORTING",
        "display": "Reporting & Analytics",
        "type": "functional",
        "rationale": "Executive summary report export and emission analytics."
    },
    {
        "id": "FR-020",
        "statement": "The system should generate carrier demurrage charge analytics and freight invoice penalty reconciliation summaries.",
        "summary": "Generate carrier demurrage analytics and penalty reconciliation summaries.",
        "canonical": "REPORTING",
        "display": "Reporting & Analytics",
        "type": "functional",
        "rationale": "Financial performance reporting and demurrage analytics."
    },
]

# 2. Build specification.docx
doc = docx.Document()
doc.add_heading("Warehouse Logistics & Fleet Dispatch System", 0)
doc.add_heading("1. System Overview", level=1)
doc.add_paragraph("The organization requires a centralized Warehouse Logistics & Fleet Dispatch System to coordinate dock scheduling, pallet inventory, driver notifications, and operational reporting.")

sections = [
    ("2. User Authentication and Terminal Access", requirements[0:4]),
    ("3. Dock Scheduling and Route Dispatch", requirements[4:8]),
    ("4. Automated Driver and Warehouse Notifications", requirements[8:12]),
    ("5. Pallet Inventory and Data Management", requirements[12:16]),
    ("6. Operational Performance Analytics and Reporting", requirements[16:20]),
]

for sec_title, req_group in sections:
    doc.add_heading(sec_title, level=1)
    for r in req_group:
        doc.add_paragraph(r["statement"])

doc_path = bench_dir / "specification.docx"
doc.save(str(doc_path))

# Also write specification.txt
txt_path = bench_dir / "specification.txt"
txt_lines = ["Warehouse Logistics & Fleet Dispatch System\n"]
for sec_title, req_group in sections:
    txt_lines.append(f"\n{sec_title}\n")
    for r in req_group:
        txt_lines.append(f"- {r['statement']}\n")
txt_path.write_text("".join(txt_lines), encoding="utf-8")

# 3. Build gold_reference.json
gold_data = {
    "project_id": "held-out-logistics-v1",
    "project_name": "Warehouse Logistics & Fleet Dispatch System",
    "domain": "Logistics & Supply Chain",
    "version": "1.0.0-held-out",
    "date_frozen": "2026-10-13",
    "source_document": "specification.docx",
    "total_requirements_evaluated": len(requirements),
    "expected_module_assignments": {r["id"]: r["canonical"] for r in requirements},
    "requirement_details": {r["id"]: r for r in requirements},
    "structural_expectations": {
        "required_capabilities": ["AUTHENTICATION", "BOOKING", "NOTIFICATION", "DATA", "REPORTING"]
    }
}

gold_path = bench_dir / "gold_reference.json"
gold_path.write_text(json.dumps(gold_data, indent=2), encoding="utf-8")

print(f"Successfully generated held-out benchmark at {bench_dir}: {len(requirements)} requirements.")
